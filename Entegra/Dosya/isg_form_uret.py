# -*- coding: utf-8 -*-
"""
KoBo (GenForms) gonderimlerini "isg yeni form.xlsx" sablonuna doldurur.

  python isg_form_uret.py --token <API_KEY> [--cikti KLASOR] [--id 1]

Cikti varsayilani: <proje>/GenForms  (ekler: GenForms/_ekler)

Her gonderim icin sablonun BIR KOPYASI uretilir; sablon degistirilmez.
Hucre eslesmesi isg_form_eslesme.json'da - yanlis hucre gorursen ORAYI duzelt,
betige dokunma.

Notlar
  - Birlesik hucrede sol-ust hucre yazilir (openpyxl kurali); bicim korunur.
  - Kod degerleri ('lisansustu', '0_rh_pozitif') KoBo'nun etiketleriyle
    degistirilmez - form tanimindaki etiketler istenirse ETIKET sozlugune eklenir.
"""
import argparse, json, os, re, sys, unicodedata
from datetime import datetime

try:
    import requests
except ImportError:
    sys.exit("requests gerekli:  pip install requests")
try:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XlImage
except ImportError:
    sys.exit("openpyxl gerekli:  pip install openpyxl")

SUNUCU  = "https://genforms-gf.genyazilim.com"
FORM_ID = "apyZkuLcNyw28bDft7nVvE"          # Ise Giris / Periyodik Muayene Formu
BURASI  = os.path.dirname(os.path.abspath(__file__))
# Tum ciktilar GenForms klasorune (proje kokunun altinda) yazilir.
CIKTI_KOK = os.path.join(os.path.dirname(BURASI), "GenForms")
SABLON  = os.path.join(BURASI, "isg yeni form.xlsx")
ESLESME = os.path.join(BURASI, "isg_form_eslesme.json")

# KoBo kod degeri -> forma yazilacak metin. Eksik olan AYNEN yazilir.
ETIKET = {
    "erkek": "Erkek", "kadin": "Kadın",
    "evli": "Evli", "bekar": "Bekâr",
    "ilkokul": "İlkokul", "ortaokul": "Ortaokul", "lise": "Lise",
    "onlisans": "Ön Lisans", "lisans": "Lisans", "lisansustu": "Lisansüstü",
    "0_rh_pozitif": "0 Rh (+)", "0_rh_negatif": "0 Rh (−)",
    "a_rh_pozitif": "A Rh (+)", "a_rh_negatif": "A Rh (−)",
    "b_rh_pozitif": "B Rh (+)", "b_rh_negatif": "B Rh (−)",
    "ab_rh_pozitif": "AB Rh (+)", "ab_rh_negatif": "AB Rh (−)",
    "periyodik": "Periyodik", "ise_giris": "İşe Giriş",
    "calisabilir": "Çalışabilir", "calisamaz": "Çalışamaz",
    "elverisli": "Elverişlidir", "elverisli_degil": "Elverişli değildir",
    "birakmis": "Bırakmış", "almiyor": "Almıyor", "icmiyor": "İçmiyor",
    "evet": "Evet", "hayir": "Hayır",
}


def etiketle(v):
    if v is None:
        return ""
    if isinstance(v, (list, dict)):
        return ""
    s = str(v).strip()
    return ETIKET.get(s, s)


def dosya_adi(s):
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9_.-]+", "_", s).strip("_")
    return s or "form"


def ek_indir(sub, token, klasor):
    """Gonderimin eklerini indirir. Doner: {question_xpath: yerel_dosya}."""
    os.makedirs(klasor, exist_ok=True)
    bulunan = {}
    for ek in sub.get("_attachments", []) or []:
        soru = ek.get("question_xpath") or ""
        url  = ek.get("download_url")
        if not soru or not url:
            continue
        ad = os.path.basename(ek.get("filename", "")) or f"{soru.replace('/', '_')}.png"
        yol = os.path.join(klasor, f"{sub.get('_id')}_{dosya_adi(ad)}")
        if not os.path.exists(yol):
            r = requests.get(url, headers={"Authorization": f"Token {token}"}, timeout=60)
            if r.status_code != 200:
                print(f"   ! ek indirilemedi ({r.status_code}): {soru}")
                continue
            with open(yol, "wb") as f:
                f.write(r.content)
        bulunan[soru] = yol
    return bulunan


def gonderimleri_al(token, tek_id=None):
    url = f"{SUNUCU}/api/v2/assets/{FORM_ID}/data/?format=json"
    r = requests.get(url, headers={"Authorization": f"Token {token}"}, timeout=60)
    r.raise_for_status()
    sonuc = r.json().get("results", [])
    if tek_id is not None:
        sonuc = [s for s in sonuc if str(s.get("_id")) == str(tek_id)]
    return sonuc


def doldur(sub, esl, cikti_klasor, ekler):
    wb = load_workbook(SABLON)
    ws = wb.worksheets[0]

    # --- duz alanlar ---
    for alan, hucre in esl["duz"].items():
        if alan in sub:
            ws[hucre] = etiketle(sub[alan])

    # --- tek birlesik bloga alt alta yazilanlar ---
    blok = esl.get("_birlesik_metin")
    if blok:
        satirlar = []
        for baslik, alan in blok["alanlar"]:
            if alan in sub:
                satirlar.append(f"{baslik}: {etiketle(sub[alan])}")
        if satirlar:
            ws[blok["hucre"]] = "\n".join(satirlar)

    # --- Hayir/Evet kutucuklari ---
    for alan, kutular in esl["kutu"].items():
        if alan.startswith("_") or alan not in sub:
            continue
        deger = str(sub[alan]).strip().lower()
        hedef = kutular.get(deger)
        if hedef:
            ws[hedef] = "X"

    # --- calisma durumu satirlari ---
    for alan, tanim in esl.get("calisma_durumu", {}).items():
        if alan.startswith("_") or alan not in sub:
            continue
        satir = tanim["satir"]
        deger = str(sub[alan]).strip().lower()
        if deger == "calisabilir":
            ws[f"H{satir}"] = "X"
        elif deger == "calisamaz":
            ws[f"J{satir}"] = "X"

    # --- onceki isler (tekrarli grup) -> 1./2./3. satirlari ---
    onceki = sub.get("onceki_isler") or []
    for i, kayit in enumerate(onceki[:3]):
        parcalar = [str(v).strip() for k, v in kayit.items()
                    if not k.startswith("_") and str(v).strip()]
        if parcalar:
            ws[f"A{20 + i}"] = f"{i + 1}. " + " / ".join(parcalar)

    # --- hekim kimlik blogu (etiket korunur, altina yazilir) ---
    hb = esl.get("_hekim_blok")
    if hb:
        satirlar = [hb.get("etiket", "")]
        for baslik, alan in hb["alanlar"]:
            if alan in sub and str(sub[alan]).strip():
                v = etiketle(sub[alan])
                satirlar.append(f"{baslik}: {v}" if baslik else v)
        ws[hb["hucre"]] = "\n".join([s for s in satirlar if s])

    # --- imza gorselleri ---
    #   Excel'de resim HUCREYE BAGLI degil, hucrenin sol-ust kosesine
    #   sabitlenir; satir yuksekligi degisirse kayabilir.
    for soru, tanim in esl.get("imza", {}).items():
        if soru.startswith("_"):
            continue
        yol = ekler.get(soru)
        if not yol or not os.path.exists(yol):
            continue
        try:
            img = XlImage(yol)
            img.width  = tanim.get("genislik", 140)
            img.height = tanim.get("yukseklik", 50)
            ws.add_image(img, tanim["hucre"])
        except Exception as e:
            print(f"   ! imza eklenemedi ({soru}): {e}")

    # Dosya adinda GONDERIM ID'si SART: ayni kisi ayni gun birden fazla form
    #   doldurabilir, ustelik Windows harf duyarsizdir - "Murat_Ozkasap" ile
    #   "murat_ozkasap" ayni dosyadir ve ikincisi birincisini EZER.
    ad = etiketle(sub.get("calisan/calisan_ad_soyad", "")) or "isimsiz"
    trh = (sub.get("hekim/muayene_tarihi") or sub.get("today")
           or datetime.now().strftime("%Y-%m-%d"))
    yol = os.path.join(
        cikti_klasor,
        f"{dosya_adi(ad)}_{dosya_adi(str(trh))}_{sub.get('_id')}.xlsx")
    wb.save(yol)
    return yol


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--token", required=True, help="KoBo API key")
    ap.add_argument("--cikti", default=CIKTI_KOK)   # varsayilan: <proje>/GenForms
    ap.add_argument("--id", default=None, help="yalniz bu gonderim (_id)")
    a = ap.parse_args()

    if not os.path.exists(SABLON):
        sys.exit(f"Sablon yok: {SABLON}")
    with open(ESLESME, encoding="utf-8") as f:
        esl = json.load(f)
    os.makedirs(a.cikti, exist_ok=True)

    subs = gonderimleri_al(a.token, a.id)
    if not subs:
        sys.exit("Gonderim bulunamadi.")
    ek_klasor = os.path.join(a.cikti, "_ekler")
    for s in subs:
        ekler = ek_indir(s, a.token, ek_klasor)
        yol = doldur(s, esl, a.cikti, ekler)
        print(f"_id={s.get('_id')}  ->  {yol}")
    print(f"\n{len(subs)} form uretildi.")


if __name__ == "__main__":
    main()
