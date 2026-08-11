# -*- coding: utf-8 -*-
"""
KoBo (GenForms) gonderimlerini RESMI Ek-2 sablonuna doldurur.

  python isg2_form_uret.py --token <API_KEY> [--id 12] [--cikti KLASOR]

Sablon : GenForms/İşe Giriş ve Muayene.xlsx      (elle doldurulmus ornekle ayni duzen)
Esleme : Dosya/isg2_eslesme.json                 (alan -> hucre; yanlissa ORAYI duzelt)
Cikti  : GenForms/  (ekler: GenForms/_ekler)     - sablonun KOPYASI uretilir

Eski surum (isg_form_uret.py) farkli bir sablona (Dosya/isg yeni form.xlsx) yaziyordu;
bu betik onun yerine gecer, eskisi bozulmaz.
"""
import argparse, json, os, re, sys, unicodedata
from datetime import datetime

try:
    import requests
except ImportError:
    sys.exit("requests gerekli:  pip install requests")
try:
    from openpyxl import load_workbook
    from openpyxl.drawing.image import Image as XlImage
except ImportError:
    sys.exit("openpyxl gerekli:  pip install openpyxl")

SUNUCU  = "https://genforms-gf.genyazilim.com"
FORM_ID = "aGT9VazSAoFSmHWx65pWZ9"          # "İşe Giriş ve Muayene" (Ek-2 duzeni, surum 3 / tek dil)
# Eski form: apyZkuLcNyw28bDft7nVvE (İşe Giriş / Periyodik Muayene Formu) -> --form ile verilebilir
BURASI  = os.path.dirname(os.path.abspath(__file__))
PROJE   = os.path.dirname(BURASI)
GENFORMS = os.path.join(PROJE, "GenForms")
SABLON  = os.path.join(GENFORMS, "İşe Giriş ve Muayene.xlsx")
ESLESME = os.path.join(BURASI, "isg2_eslesme.json")

# KoBo kod degeri -> forma yazilacak metin. Listede olmayan AYNEN yazilir.
ETIKET = {
    "evet": "Evet", "hayir": "Hayır",
    "erkek": "Erkek", "kadin": "Kadın",
    "evli": "Evli", "bekar": "Bekâr", "diger": "Diğer",
    "okuryazar_degil": "Okur-yazar değil", "ilkokul": "İlkokul",
    "ortaokul": "Ortaokul", "lise": "Lise", "onlisans": "Ön lisans",
    "lisans": "Lisans", "lisansustu": "Lisansüstü",
    "0_rh_pozitif": "0 Rh +", "0_rh_negatif": "0 Rh −",
    "a_rh_pozitif": "A Rh +", "a_rh_negatif": "A Rh −",
    "b_rh_pozitif": "B Rh +", "b_rh_negatif": "B Rh −",
    "ab_rh_pozitif": "AB Rh +", "ab_rh_negatif": "AB Rh −",
    "bilinmiyor": "Bilinmiyor",
    "icmiyor": "İçmiyor", "iciyor": "İçiyor", "birakmis": "Bırakmış",
    "almiyor": "Almıyor", "aliyor": "Alıyor",
    "calisabilir": "Çalışabilir", "calisamaz": "Çalışamaz",
    "ise_giris": "İşe giriş", "periyodik": "Periyodik muayene",
}


ISO_TARIH = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


def tarih_bicim(s):
    """KoBo tarihleri ISO gelir (2026-08-07). Formda GUN/AY/YIL yazilir."""
    m = ISO_TARIH.match(s)
    return "%s/%s/%s" % (m.group(3), m.group(2), m.group(1)) if m else s


def etiketle(v):
    if v is None or isinstance(v, (list, dict)):
        return ""
    s = str(v).strip()
    s = ETIKET.get(s, s)
    return tarih_bicim(s)


def dosya_adi(s):
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9_.-]+", "_", s).strip("_")
    return s or "form"


def ek_indir(sub, token, klasor):
    """Gonderimin eklerini indirir. Doner: {question_xpath: yerel_dosya}."""
    os.makedirs(klasor, exist_ok=True)
    bulunan = {}
    for ek in sub.get("_attachments", []) or []:
        soru = ek.get("question_xpath") or ""
        url = ek.get("download_url")
        if not soru or not url:
            continue
        ad = os.path.basename(ek.get("filename", "")) or (soru.replace("/", "_") + ".png")
        yol = os.path.join(klasor, "%s_%s" % (sub.get("_id"), dosya_adi(ad)))
        if not os.path.exists(yol):
            r = requests.get(url, headers={"Authorization": "Token " + token}, timeout=60)
            if r.status_code != 200:
                print("   ! ek indirilemedi (%s): %s" % (r.status_code, soru))
                continue
            with open(yol, "wb") as f:
                f.write(r.content)
        bulunan[soru] = yol
    return bulunan


def gonderimleri_al(token, tek_id=None, form_id=None):
    url = "%s/api/v2/assets/%s/data/?format=json" % (SUNUCU, form_id or FORM_ID)
    r = requests.get(url, headers={"Authorization": "Token " + token}, timeout=60)
    r.raise_for_status()
    sonuc = r.json().get("results", [])
    if tek_id is not None:
        sonuc = [s for s in sonuc if str(s.get("_id")) == str(tek_id)]
    return sonuc


def _al(sub, alan):
    """Gonderimden alan degeri: duz anahtar ya da grup icinde."""
    if alan in sub:
        return sub[alan]
    # bazi sunumlarda gruplar ic ice sozluk olarak gelir
    parcalar = alan.split("/")
    dugum = sub
    for p in parcalar:
        if isinstance(dugum, dict) and p in dugum:
            dugum = dugum[p]
        else:
            return None
    return dugum


def doldur(sub, esl, cikti_klasor, ekler):
    wb = load_workbook(SABLON)
    ws = wb[esl["sayfa"]] if esl.get("sayfa") in wb.sheetnames else wb.worksheets[0]

    # --- duz alanlar ---
    for alan, hucre in esl.get("duz", {}).items():
        if alan.startswith("_"):
            continue
        v = _al(sub, alan)
        if v not in (None, ""):
            ws[hucre] = etiketle(v)

    # --- X konan kutucuklar ---
    for alan, kutular in esl.get("kutu", {}).items():
        if alan.startswith("_"):
            continue
        v = _al(sub, alan)
        if v in (None, ""):
            continue
        hedef = kutular.get(str(v).strip().lower())
        if hedef:
            ws[hedef] = "X"

    # --- etiketli kutucuklar: hucrede "Hayır"/"Evet" YAZISI var; ustune X yazmak
    #     etiketi siler. Secilen ☒, digeri ☐ ile isaretlenir, etiket korunur.
    for alan, kutular in esl.get("isaret", {}).items():
        if alan.startswith("_"):
            continue
        v = _al(sub, alan)
        if v in (None, ""):
            continue
        secilen = str(v).strip().lower()
        for deger, hucre in kutular.items():
            if not hucre:
                continue
            mevcut = ws[hucre].value
            etiket = str(mevcut).strip() if mevcut not in (None, "") else ETIKET.get(deger, deger)
            etiket = etiket.lstrip("☒☐ ").strip()
            ws[hucre] = ("☒ " if deger == secilen else "☐ ") + etiket

    # --- tek hucreye METIN yazan kutular (EVET/HAYIR) ---
    for alan, tanim in esl.get("metin_kutu", {}).items():
        if alan.startswith("_"):
            continue
        v = _al(sub, alan)
        if v in (None, ""):
            continue
        metin = tanim.get("metin", {}).get(str(v).strip().lower())
        if metin:
            ws[tanim["hucre"]] = metin

    # --- birden fazla alan tek hucrede ---
    for hucre, tanim in esl.get("birlesik", {}).items():
        if hucre.startswith("_"):
            continue
        parcalar = [etiketle(_al(sub, a)) for a in tanim["alanlar"]]
        parcalar = [p for p in parcalar if p]
        if parcalar:
            ws[hucre] = tanim.get("ayrac", " ").join(parcalar)

    # --- etiketi korunan hucreler ---
    for hucre, tanim in esl.get("onekli", {}).items():
        if hucre.startswith("_"):
            continue
        v = etiketle(_al(sub, tanim["alan"]))
        if not v:
            continue
        if tanim.get("sablon"):
            ws[hucre] = tanim["sablon"].replace("{deger}", v)
        else:
            ws[hucre] = tanim.get("onek", "") + v

    # --- daha once calistigi isyeri (tek satir) ---
    ts = esl.get("tekil_satir")
    if ts:
        for alan, sutun in ts.get("sutunlar", {}).items():
            if alan.startswith("_"):
                continue
            v = etiketle(_al(sub, alan))
            if v:
                ws["%s%d" % (sutun, ts["satir"])] = v

    # --- hekim blogu: etiket korunur, bilgiler altina ---
    hb = esl.get("hekim_blok")
    if hb:
        satirlar = [hb.get("etiket", "")]
        for baslik, alan in hb["alanlar"]:
            v = etiketle(_al(sub, alan))
            if v:
                satirlar.append(("%s: %s" % (baslik, v)) if baslik else v)
        ws[hb["hucre"]] = "\n".join([s for s in satirlar if s])

    # --- gorseller (fotograf / imza) ---
    #   Excel'de resim hucreye BAGLI degil, hucrenin sol-ust kosesine sabitlenir.
    for soru, tanim in esl.get("imza", {}).items():
        if soru.startswith("_"):
            continue
        yol = ekler.get(soru)
        if not yol or not os.path.exists(yol):
            continue
        try:
            img = XlImage(yol)
            img.width = tanim.get("genislik", 140)
            img.height = tanim.get("yukseklik", 50)
            ws.add_image(img, tanim["hucre"])
        except Exception as e:
            print("   ! gorsel eklenemedi (%s): %s" % (soru, e))

    # Dosya adinda GONDERIM ID'si SART: ayni kisi ayni gun birden fazla form
    #   doldurabilir; Windows harf duyarsiz oldugu icin ikincisi birincisini EZER.
    ad = etiketle(_al(sub, "calisan/calisan_ad_soyad")) or "isimsiz"
    trh = (_al(sub, "hekim/muayene_tarihi") or sub.get("today")
           or datetime.now().strftime("%Y-%m-%d"))
    yol = os.path.join(cikti_klasor,
                       "%s_%s_%s.xlsx" % (dosya_adi(ad), dosya_adi(str(trh)), sub.get("_id")))
    wb.save(yol)
    return yol


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--token", required=True, help="KoBo API key")
    ap.add_argument("--cikti", default=GENFORMS, help="cikti klasoru (varsayilan: GenForms)")
    ap.add_argument("--id", default=None, help="yalniz bu gonderim (_id)")
    ap.add_argument("--form", default=FORM_ID, help="form uid (varsayilan: yeni Ek-2 formu)")
    a = ap.parse_args()

    if not os.path.exists(SABLON):
        sys.exit("Sablon yok: " + SABLON)
    with open(ESLESME, encoding="utf-8") as f:
        esl = json.load(f)
    os.makedirs(a.cikti, exist_ok=True)

    subs = gonderimleri_al(a.token, a.id, a.form)
    if not subs:
        sys.exit("Gonderim bulunamadi.")
    ek_klasor = os.path.join(a.cikti, "_ekler")
    for s in subs:
        ekler = ek_indir(s, a.token, ek_klasor)
        yol = doldur(s, esl, a.cikti, ekler)
        print("_id=%s  ->  %s" % (s.get("_id"), yol))
    print("\n%d form uretildi." % len(subs))


if __name__ == "__main__":
    main()
